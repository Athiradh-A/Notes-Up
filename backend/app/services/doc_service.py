import io
import os
import re
from pypdf import PdfReader
from pptx import Presentation
from fastapi import UploadFile
from typing import List, Dict

def sanitize_filename(filename: str) -> str:
    """Removes potentially dangerous characters from filenames."""
    return re.sub(r'[^a-zA-Z0-9._-]', '_', os.path.basename(filename))

async def extract_text(file: UploadFile) -> List[Dict]:
    """
    Extracts text from PDF or PPTX and returns a list of content per page/slide.
    Returns: List of {'page': int, 'content': str}
    """
    pages_content = []
    try:
        # Sanitize filename and validate extension
        filename = sanitize_filename(file.filename)
        file_ext = filename.split('.')[-1].lower()
        
        if file_ext not in ['pdf', 'pptx']:
            raise ValueError(f"Unsupported file extension: {file_ext}")

        file_bytes = await file.read()
        if file_ext == 'pdf':
            reader = PdfReader(io.BytesIO(file_bytes))
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages_content.append({'page': i + 1, 'content': text})
        elif file_ext == 'pptx':
            prs = Presentation(io.BytesIO(file_bytes))
            for i, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text.append(shape.text)
                full_text = " ".join(slide_text)
                if full_text.strip():
                    pages_content.append({'page': i + 1, 'content': full_text})
    except Exception as e:
        print(f"Error extracting text from {file.filename}: {e}")
        raise e
    return pages_content
